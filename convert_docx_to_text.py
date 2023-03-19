import difflib
import re

from docx import Document
from underthesea import sent_tokenize, text_normalize

def convert(filename):
    document = Document(filename)
    try:
        sentences = []
        for para in document.paragraphs:
            for sentence in sent_tokenize(text_normalize(para.text)):
                sentences.append(
                    sentence
                )
        tables = []
        for table in document.tables:
            for row in table.rows:
                cells = []
                for cell in row.cells:
                    cells.append(cell.text.replace('\n', ' '))
                tables.append(str('-'.join(cells)))
        return sentences, tables
    except Exception as e:
        return e
    
def tokenize(s):
    return re.split('\s+', s)
def untokenize(ts):
    return ' '.join(ts)
        
def equalize(s1, s2):
    l1 = tokenize(s1)
    l2 = tokenize(s2)
    res1 = []
    res2 = []
    prev = difflib.Match(0,0,0)
    for match in difflib.SequenceMatcher(a=l1, b=l2).get_matching_blocks():
        if (prev.a + prev.size != match.a):
            for i in range(prev.a + prev.size, match.a):
                res2 += ['_' * len(l1[i])]
            res1 += l1[prev.a + prev.size:match.a]
        if (prev.b + prev.size != match.b):
            for i in range(prev.b + prev.size, match.b):
                res1 += ['_' * len(l2[i])]
            res2 += l2[prev.b + prev.size:match.b]
        res1 += l1[match.a:match.a+match.size]
        res2 += l2[match.b:match.b+match.size]
        prev = match
    return untokenize(res1), untokenize(res2)

def insert_newlines(string, every=64, window=10):
    result = []
    from_string = string
    while len(from_string) > 0:
        cut_off = every
        if len(from_string) > every:
            while (from_string[cut_off-1] != ' ') and (cut_off > (every-window)):
                cut_off -= 1
        else:
            cut_off = len(from_string)
        part = from_string[:cut_off]
        result += [part]
        from_string = from_string[cut_off:]
    return result

def show_comparison(s1, s2, width=40, margin=10, sidebyside=True, compact=False):
    s1, s2 = equalize(s1,s2)

    if sidebyside:
        s1 = insert_newlines(s1, width, margin)
        s2 = insert_newlines(s2, width, margin)
        if compact:
            for i in range(0, len(s1)):
                lft = re.sub(' +', ' ', s1[i].replace('_', '')).ljust(width)
                rgt = re.sub(' +', ' ', s2[i].replace('_', '')).ljust(width) 
                print(lft + ' | ' + rgt + ' | ')        
        else:
            for i in range(0, len(s1)):
                lft = s1[i].ljust(width)
                rgt = s2[i].ljust(width)
                print(lft + ' | ' + rgt + ' | ')
    else:
        print(s1)
        print(s2)

if __name__=='__main__':
    s1, tables_ = convert("tai_lieu.docx")
    s2, tables = convert("tai_lieu.docx")
    show_comparison(' '.join(tables_), ' '.join(tables_), sidebyside=False)